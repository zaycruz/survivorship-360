"""Build deck.pptx: one full-bleed slide PNG per slide plus speaker notes from slides.py.

Run: uv run --with python-pptx python -I build_pptx.py
"""
import os
import shutil
import sys

from pptx import Presentation
from pptx.util import Inches

here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, here)
from slides import SLIDES  # noqa: E402

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]

for i, s in enumerate(SLIDES, 1):
    png = os.path.join(here, "slides", f"{i:02d}.png")
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(png, 0, 0, prs.slide_width, prs.slide_height)
    slide.notes_slide.notes_text_frame.text = s["notes"]

out = os.path.join(here, "deck.pptx")
prs.save(out)
dl = os.path.expanduser("~/Downloads")
shutil.copy(out, os.path.join(dl, "survivorship-360-pitch.pptx"))
shutil.copy(os.path.join(here, "deck.pdf"), os.path.join(dl, "survivorship-360-pitch.pdf"))
print(f"wrote {out} ({len(SLIDES)} slides) and copies in ~/Downloads")
