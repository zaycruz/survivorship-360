"""E2E check: the exported deck has one full-bleed image and notes per slide, and every number traces to its module."""
import os
import sys

from pptx import Presentation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from slides import CLAIMS, ROOT, SLIDES  # noqa: E402

here = os.path.dirname(os.path.abspath(__file__))
lines, ok = [], True


def check(cond, msg):
    global ok
    ok &= bool(cond)
    lines.append(("PASS " if cond else "FAIL ") + msg)


for path in [os.path.join(here, "deck.pptx"), os.path.expanduser("~/Downloads/survivorship-360-pitch.pptx")]:
    prs = Presentation(path)
    check(len(prs.slides) == len(SLIDES), f"{os.path.basename(path)}: {len(prs.slides)} slides, expected {len(SLIDES)}")
    for i, s in enumerate(prs.slides, 1):
        pics = [sh for sh in s.shapes if sh.shape_type == 13]
        full = pics and pics[0].width == prs.slide_width and pics[0].height == prs.slide_height
        notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
        check(full, f"  slide {i}: full-bleed image")
        check(len(notes) > 40, f"  slide {i}: notes ({len(notes)} chars)")

for f in ["deck.pdf", os.path.expanduser("~/Downloads/survivorship-360-pitch.pdf")]:
    p = os.path.join(here, f)
    check(os.path.getsize(p) > 100_000, f"{os.path.basename(p)} exists ({os.path.getsize(p)} bytes)")

pngs = sorted(x for x in os.listdir(os.path.join(here, "slides")) if x.endswith(".png"))
check(len(pngs) == len(SLIDES), f"slides/: {len(pngs)} PNGs")

lines.append("Numbers traced to module docs (read now):")
for lit, mod, fn in CLAIMS:
    text = open(os.path.join(os.path.expanduser(ROOT), mod, fn), encoding="utf-8").read()
    check(lit in text, f"  {mod}/{fn}: {lit!r}")

lines.append("RESULT: " + ("PASS" if ok else "FAIL"))
out = "\n".join(lines)
open(os.path.join(here, "check.txt"), "w").write(out + "\n")
print(out)
sys.exit(0 if ok else 1)
